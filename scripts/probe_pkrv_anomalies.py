"""Probe PKRV anomaly artifacts without saving production files."""

from __future__ import annotations

import csv
import hashlib
from collections import defaultdict
from io import BytesIO, StringIO
from pathlib import Path

import requests
from openpyxl import load_workbook

from pakyield.config import PKRV_CANONICAL_TENORS

PROJECT_ROOT = Path(__file__).resolve().parents[1]

ANOMALY_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_bulk_anomalies.csv"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/153.0.0.0 Safari/537.36"
    ),
    "Referer": (
        "https://www.mufap.com.pk/"
        "WebRegulations/Index"
        "?Head=Pricing"
        "&title=PKRV%2FPKISRV%2FPKFRV"
    ),
}


def sha256_bytes(content: bytes) -> str:
    """Return SHA-256 for in-memory bytes."""

    return hashlib.sha256(content).hexdigest()


def inspect_csv(content: bytes) -> dict[str, object]:
    """Inspect a CSV artifact without saving it."""

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return {
            "structure": "CSV_DECODE_FAILED",
            "rows": None,
            "header": None,
            "tenors": None,
        }

    rows = list(csv.reader(StringIO(text)))

    if not rows:
        return {
            "structure": "EMPTY_CSV",
            "rows": 0,
            "header": None,
            "tenors": None,
        }

    header = rows[0]

    tenors = tuple(row[0].strip() for row in rows[1:] if row)

    canonical = tenors == PKRV_CANONICAL_TENORS

    return {
        "structure": ("CANONICAL_PKRV_CSV" if canonical else "NONCANONICAL_CSV"),
        "rows": len(rows),
        "header": header,
        "tenors": tenors,
    }


def inspect_xlsx(content: bytes) -> dict[str, object]:
    """Inspect an XLSX workbook in memory."""

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        return {
            "structure": (f"XLSX_READ_FAILED:{type(exc).__name__}"),
            "sheets": None,
            "preview": None,
        }

    sheets = workbook.sheetnames

    preview: list[list[object]] = []

    if sheets:
        sheet = workbook[sheets[0]]

        for row in sheet.iter_rows(
            min_row=1,
            max_row=25,
            values_only=True,
        ):
            preview.append(list(row))

    workbook.close()

    return {
        "structure": "XLSX_READABLE",
        "sheets": sheets,
        "preview": preview,
    }


def inspect_bytes(
    file_name: str,
    content: bytes,
) -> dict[str, object]:
    """Inspect content according to extension."""

    suffix = Path(file_name).suffix.lower()

    if suffix == ".csv":
        return inspect_csv(content)

    if suffix == ".xlsx":
        return inspect_xlsx(content)

    if suffix == ".xls":
        return {
            "structure": "XLS_BINARY",
            "magic_hex": content[:16].hex(),
            "preview": None,
        }

    return {
        "structure": "UNKNOWN_EXTENSION",
    }


def main() -> None:
    """Probe all REVIEW artifacts without storing them."""

    with ANOMALY_PATH.open(
        newline="",
        encoding="utf-8",
    ) as file:
        rows = list(csv.DictReader(file))

    assert len(rows) == 13

    by_date: dict[
        str,
        list[dict[str, object]],
    ] = defaultdict(list)

    with requests.Session() as session:
        for index, row in enumerate(
            rows,
            start=1,
        ):
            url = row["source_url"]
            file_name = row["file_name"]

            response = session.get(
                url,
                headers=HEADERS,
                timeout=30,
            )

            status = response.status_code
            content = response.content

            result = {
                "file_name": file_name,
                "status": status,
                "content_type": (
                    response.headers.get(
                        "Content-Type",
                        "",
                    )
                ),
                "size": len(content),
                "sha256": (sha256_bytes(content) if content else ""),
            }

            if status == 200 and content:
                result.update(
                    inspect_bytes(
                        file_name,
                        content,
                    )
                )

            by_date[row["observation_date"]].append(result)

            print(
                f"[{index:02d}/13]",
                row["observation_date"],
                "|",
                file_name,
                "| HTTP",
                status,
                "| bytes",
                len(content),
                "|",
                result.get(
                    "structure",
                    "NOT_INSPECTED",
                ),
            )

    print("")
    print("===== DETAILED ANOMALY PROBE =====")

    for observation_date in sorted(by_date):
        print("")
        print(
            "DATE:",
            observation_date,
        )

        items = by_date[observation_date]

        for item in items:
            print(
                "  FILE:",
                item["file_name"],
            )
            print(
                "    HTTP:",
                item["status"],
            )
            print(
                "    TYPE:",
                item["content_type"],
            )
            print(
                "    SIZE:",
                item["size"],
            )
            print(
                "    SHA256:",
                item["sha256"],
            )
            print(
                "    STRUCTURE:",
                item.get("structure"),
            )

            header = item.get("header")

            if header is not None:
                print(
                    "    HEADER:",
                    header,
                )

            tenors = item.get("tenors")

            if tenors is not None:
                print(
                    "    TENORS:",
                    tenors,
                )

            sheets = item.get("sheets")

            if sheets is not None:
                print(
                    "    SHEETS:",
                    sheets,
                )

            preview = item.get("preview")

            if preview is not None:
                print("    PREVIEW:")

                for row in preview[:12]:
                    print(
                        "     ",
                        row,
                    )

            magic_hex = item.get("magic_hex")

            if magic_hex is not None:
                print(
                    "    MAGIC:",
                    magic_hex,
                )

        if len(items) > 1:
            hashes = {str(item["sha256"]) for item in items}

            print(
                "  DUPLICATE FILES BYTE-IDENTICAL:",
                len(hashes) == 1,
            )

    print("")
    print("PASS: anomaly probe completed without saving raw artifacts")


if __name__ == "__main__":
    main()
