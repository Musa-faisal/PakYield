"""Build structural census for the frozen MUFAP PKISRV raw archive."""

from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]

PLAN_PATH = PROJECT_ROOT / "data" / "manifests" / "pkisrv_acquisition_plan.csv"

RAW_DIR = PROJECT_ROOT / "data" / "raw" / "mufap" / "pkisrv"

OUTPUT_PATH = PROJECT_ROOT / "data" / "manifests" / "pkisrv_structural_census.csv"

EXPECTED_PLAN_ROWS = 407

EXPECTED_TENORS = (
    "1M",
    "3M",
    "6M",
    "9M",
    "1Y",
)

EXPECTED_TENOR_SET = set(EXPECTED_TENORS)

OUTPUT_FIELDS = [
    "observation_date",
    "file_name",
    "extension",
    "detected_format",
    "detected_encoding",
    "tenor_count",
    "found_tenors",
    "missing_tenors",
    "duplicate_tenors",
    "unparseable_tenors",
    "structural_class",
    "validation_status",
]


def decode_text_bytes(raw: bytes) -> tuple[str, str]:
    """Decode text-like MUFAP artifacts conservatively."""

    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "utf-16"

    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig"), "utf-8-sig"

    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1252"), "cp1252"


def normalize_tenor(value: Any) -> str | None:
    """Normalize standardized MUFAP PKISRV tenor labels."""

    if value is None:
        return None

    text = str(value).strip().lower()

    text = text.replace("–", "-").replace("—", "-").replace("_", " ").replace("-", " ")

    text = re.sub(
        r"\s+",
        " ",
        text,
    ).strip()

    aliases = {
        "1 month": "1M",
        "1month": "1M",
        "1 m": "1M",
        "1m": "1M",
        "3 month": "3M",
        "3month": "3M",
        "3 m": "3M",
        "3m": "3M",
        "6 month": "6M",
        "6month": "6M",
        "6 m": "6M",
        "6m": "6M",
        "9 month": "9M",
        "9month": "9M",
        "9 m": "9M",
        "9m": "9M",
        "1 year": "1Y",
        "1year": "1Y",
        "1 y": "1Y",
        "1y": "1Y",
    }

    return aliases.get(text)


def parse_rate(value: Any) -> float | None:
    """Parse a benchmark yield value without altering it."""

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):
        return float(value)

    text = str(value).strip()

    if not text:
        return None

    text = text.replace("%", "").replace(",", "").strip()

    try:
        return float(text)
    except ValueError:
        return None


def scan_rows(
    rows: list[list[Any]],
) -> tuple[
    dict[str, float | None],
    set[str],
    set[str],
]:
    """Extract standardized tenor labels and adjacent values."""

    found: dict[str, float | None] = {}
    duplicates: set[str] = set()
    unparseable: set[str] = set()

    for row in rows:
        for index, value in enumerate(row):
            tenor = normalize_tenor(value)

            if tenor is None:
                continue

            if tenor in found:
                duplicates.add(tenor)
                continue

            next_value = row[index + 1] if index + 1 < len(row) else None

            rate = parse_rate(next_value)

            found[tenor] = rate

            if rate is None:
                unparseable.add(tenor)

    return (
        found,
        duplicates,
        unparseable,
    )


def rows_from_csv(
    raw: bytes,
) -> tuple[
    list[list[Any]],
    str,
    str,
]:
    """Read a CSV-like artifact."""

    text, encoding = decode_text_bytes(raw)

    rows = [list(row) for row in csv.reader(text.splitlines())]

    return (
        rows,
        "CSV_TEXT",
        encoding,
    )


def rows_from_xlsx(
    raw: bytes,
) -> tuple[
    list[list[Any]],
    str,
    str,
]:
    """Read an OOXML workbook."""

    workbook = load_workbook(
        BytesIO(raw),
        read_only=True,
        data_only=True,
    )

    rows: list[list[Any]] = []

    for sheet in workbook.worksheets:
        for values in sheet.iter_rows(values_only=True):
            rows.append(list(values))

    workbook.close()

    return (
        rows,
        "XLSX",
        "OOXML",
    )


def rows_from_xls(
    raw: bytes,
) -> tuple[
    list[list[Any]],
    str,
    str,
]:
    """Read the known Spreadsheet XML .xls representation."""

    text, encoding = decode_text_bytes(raw)

    stripped = text.lstrip()

    if not (stripped.startswith("<?xml") or "<Workbook" in stripped[:500]):
        raise RuntimeError("Unsupported binary .xls artifact")

    root = ET.fromstring(text)

    rows: list[list[Any]] = []

    for row in root.findall(".//{*}Row"):
        values: list[Any] = []

        for cell in row.findall("{*}Cell"):
            data = cell.find("{*}Data")

            values.append(data.text if data is not None else None)

        rows.append(values)

    return (
        rows,
        "SPREADSHEET_XML",
        encoding,
    )


def classify(
    found: dict[str, float | None],
    duplicates: set[str],
    unparseable: set[str],
) -> tuple[str, str]:
    """Classify one raw PKISRV artifact."""

    found_set = set(found)

    if duplicates:
        return (
            "DUPLICATE_TENOR",
            "REVIEW",
        )

    if unparseable:
        return (
            "UNPARSEABLE_RATE",
            "REVIEW",
        )

    if found_set == EXPECTED_TENOR_SET:
        return (
            "COMPLETE_5",
            "ACCEPTED",
        )

    if not found_set:
        return (
            "ABSENT",
            "REVIEW",
        )

    return (
        "PARTIAL",
        "REVIEW",
    )


def main() -> None:
    """Build and write the structural census."""

    with PLAN_PATH.open(
        newline="",
        encoding="utf-8",
    ) as file:
        plan = list(csv.DictReader(file))

    assert len(plan) == EXPECTED_PLAN_ROWS

    plan_names = {row["file_name"] for row in plan}

    assert len(plan_names) == EXPECTED_PLAN_ROWS

    physical_files = {
        path.name: path
        for path in RAW_DIR.iterdir()
        if (
            path.is_file()
            and path.name != ".gitkeep"
            and not path.name.endswith(".part")
        )
    }

    assert len(physical_files) == EXPECTED_PLAN_ROWS
    assert set(physical_files) == plan_names

    part_files = list(RAW_DIR.glob("*.part"))

    assert not part_files

    results = []

    for index, metadata in enumerate(
        plan,
        start=1,
    ):
        path = physical_files[metadata["file_name"]]

        raw = path.read_bytes()

        extension = path.suffix.lower()

        if extension == ".csv":
            (
                rows,
                detected_format,
                encoding,
            ) = rows_from_csv(raw)

        elif extension == ".xlsx":
            (
                rows,
                detected_format,
                encoding,
            ) = rows_from_xlsx(raw)

        elif extension == ".xls":
            (
                rows,
                detected_format,
                encoding,
            ) = rows_from_xls(raw)

        else:
            raise RuntimeError(f"Unexpected PKISRV extension: {extension}")

        (
            found,
            duplicates,
            unparseable,
        ) = scan_rows(rows)

        missing = EXPECTED_TENOR_SET - set(found)

        (
            structural_class,
            validation_status,
        ) = classify(
            found,
            duplicates,
            unparseable,
        )

        results.append(
            {
                "observation_date": metadata["observation_date"],
                "file_name": metadata["file_name"],
                "extension": extension,
                "detected_format": detected_format,
                "detected_encoding": encoding,
                "tenor_count": str(len(found)),
                "found_tenors": "|".join(
                    tenor for tenor in EXPECTED_TENORS if tenor in found
                ),
                "missing_tenors": "|".join(
                    tenor for tenor in EXPECTED_TENORS if tenor in missing
                ),
                "duplicate_tenors": "|".join(sorted(duplicates)),
                "unparseable_tenors": "|".join(sorted(unparseable)),
                "structural_class": structural_class,
                "validation_status": validation_status,
            }
        )

        if index % 100 == 0 or index == EXPECTED_PLAN_ROWS:
            print(f"Scanned {index}/{EXPECTED_PLAN_ROWS}")

    assert len(results) == EXPECTED_PLAN_ROWS

    temporary = OUTPUT_PATH.with_suffix(OUTPUT_PATH.suffix + ".part")

    with temporary.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=OUTPUT_FIELDS,
        )

        writer.writeheader()
        writer.writerows(results)

    temporary.replace(OUTPUT_PATH)

    classes = Counter(row["structural_class"] for row in results)

    statuses = Counter(row["validation_status"] for row in results)

    print("")
    print("===== PKISRV STRUCTURAL CENSUS =====")

    print(
        "Rows:",
        len(results),
    )

    print("")
    print("===== STRUCTURAL CLASSES =====")

    for name, count in sorted(classes.items()):
        print(
            name,
            count,
        )

    print("")
    print("===== VALIDATION STATUS =====")

    for name, count in sorted(statuses.items()):
        print(
            name,
            count,
        )

    print("")
    print("===== REVIEW RECORDS =====")

    review = [row for row in results if row["validation_status"] != "ACCEPTED"]

    print(
        "Count:",
        len(review),
    )

    for row in review:
        print(
            row["observation_date"],
            "|",
            row["file_name"],
            "|",
            row["structural_class"],
            "| found:",
            row["found_tenors"] or "<NONE>",
            "| missing:",
            row["missing_tenors"] or "<NONE>",
            "| duplicate:",
            row["duplicate_tenors"] or "<NONE>",
            "| unparseable:",
            row["unparseable_tenors"] or "<NONE>",
        )

    print("")
    print("===== ACCEPTED BY YEAR =====")

    accepted_by_year = Counter(
        row["observation_date"][:4]
        for row in results
        if row["validation_status"] == "ACCEPTED"
    )

    for year, count in sorted(accepted_by_year.items()):
        print(
            year,
            count,
        )

    print("")
    print(
        "Manifest:",
        OUTPUT_PATH.relative_to(PROJECT_ROOT),
    )

    print("")
    print("PASS: PKISRV structural census complete")


if __name__ == "__main__":
    main()
