"""Build the authoritative PKRV raw structural-acceptance manifest."""

from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from pakyield.validation.pkrv_raw import (
    CANONICAL_TENORS,
    decode_text_bytes,
    extract_canonical_tenors,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]

PLAN_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_approved_acquisition_plan.csv"

SOURCE_GAPS_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_source_gaps.csv"

EXCLUSIONS_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_validation_exclusions.csv"

RAW_DIR = PROJECT_ROOT / "data" / "raw" / "mufap" / "pkrv"

OUTPUT_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_structural_acceptance.csv"

PARTIAL_14_MISSING = {
    "1M",
    "2M",
    "3M",
    "4M",
    "6M",
    "9M",
}

FIELDS = [
    "observation_date",
    "file_name",
    "file_extension",
    "detected_encoding",
    "tenor_count",
    "missing_tenors",
    "structural_class",
    "validation_status",
]


def read_csv(
    path: Path,
) -> list[dict[str, str]]:
    """Read a project manifest."""

    with path.open(
        newline="",
        encoding="utf-8",
    ) as file:
        return list(csv.DictReader(file))


def xlsx_as_text(
    path: Path,
) -> str:
    """Render XLSX cell values to text for structural tenor detection."""

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    parts: list[str] = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                if value is not None:
                    parts.append(str(value))

    workbook.close()

    return "\n".join(parts)


def artifact_text(
    path: Path,
) -> tuple[str, str]:
    """Return artifact text and detected representation."""

    suffix = path.suffix.lower()

    if suffix == ".xlsx":
        return (
            xlsx_as_text(path),
            "xlsx-openpyxl",
        )

    if suffix in {
        ".csv",
        ".xls",
    }:
        return decode_text_bytes(path.read_bytes())

    raise RuntimeError(f"Unsupported PKRV raw extension: {path}")


def main() -> None:
    """Build and validate final PKRV structural classifications."""

    plan = read_csv(PLAN_PATH)

    source_gaps = read_csv(SOURCE_GAPS_PATH)

    exclusions = read_csv(EXCLUSIONS_PATH)

    if len(plan) != 1159:
        raise RuntimeError(
            f"Expected 1,159 approved PKRV metadata dates; found {len(plan)}"
        )

    gap_names = {row["file_name"] for row in source_gaps if row["status"] == "SKIP"}

    if len(gap_names) != 2:
        raise RuntimeError(
            f"Expected exactly two accepted source gaps; found {len(gap_names)}"
        )

    exclusion_names = {
        row["file_name"] for row in exclusions if row["status"] == "EXCLUDE"
    }

    if len(exclusion_names) != 8:
        raise RuntimeError(
            "Expected exactly eight validation exclusions; "
            f"found {len(exclusion_names)}"
        )

    if gap_names & exclusion_names:
        raise RuntimeError(
            "A file cannot be both a source gap and an acquired validation exclusion."
        )

    output: list[dict[str, str]] = []

    physically_missing: list[str] = []

    unexpected_missing: list[str] = []

    for row in plan:
        file_name = row["file_name"]
        path = RAW_DIR / file_name

        if file_name in gap_names:
            if path.exists():
                raise RuntimeError(f"Source-gap file unexpectedly exists: {file_name}")

            physically_missing.append(file_name)

            continue

        if not path.is_file():
            unexpected_missing.append(file_name)
            continue

        text, encoding = artifact_text(path)

        tenors = extract_canonical_tenors(text)

        missing = [tenor for tenor in CANONICAL_TENORS if tenor not in tenors]

        if file_name in exclusion_names:
            structural_class = "EXCLUDED_WRONG_DATASET"
            validation_status = "EXCLUDE"

        elif len(tenors) == 20:
            structural_class = "FULL_20"
            validation_status = "ACCEPT"

        elif len(tenors) == 14 and set(missing) == PARTIAL_14_MISSING:
            structural_class = "PARTIAL_14"
            validation_status = "ACCEPT"

        else:
            raise RuntimeError(
                "Unresolved PKRV structure: "
                f"{file_name} | "
                f"tenors={len(tenors)} | "
                f"missing={missing}"
            )

        output.append(
            {
                "observation_date": row["observation_date"],
                "file_name": file_name,
                "file_extension": (path.suffix.lower()),
                "detected_encoding": (encoding),
                "tenor_count": str(len(tenors)),
                "missing_tenors": "|".join(missing),
                "structural_class": (structural_class),
                "validation_status": (validation_status),
            }
        )

    if unexpected_missing:
        raise RuntimeError(
            f"Unexpected approved raw artifacts are missing: {unexpected_missing}"
        )

    if len(physically_missing) != 2:
        raise RuntimeError(
            "Expected exactly two physical source gaps; "
            f"found {len(physically_missing)}"
        )

    if len(output) != 1157:
        raise RuntimeError(
            "Expected 1,157 acquired artifacts in "
            "structural acceptance; "
            f"found {len(output)}"
        )

    classifications = Counter(row["structural_class"] for row in output)

    expected = {
        "FULL_20": 1119,
        "PARTIAL_14": 30,
        "EXCLUDED_WRONG_DATASET": 8,
    }

    if dict(classifications) != expected:
        raise RuntimeError(
            "Structural-class counts do not match "
            "the audited contract. "
            f"Found: {dict(classifications)}"
        )

    temporary = OUTPUT_PATH.with_suffix(OUTPUT_PATH.suffix + ".part")

    with temporary.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=FIELDS,
        )

        writer.writeheader()
        writer.writerows(output)

    temporary.replace(OUTPUT_PATH)

    print("===== PKRV STRUCTURAL ACCEPTANCE =====")

    print(
        "Approved metadata dates:",
        len(plan),
    )

    print(
        "Source gaps:",
        len(physically_missing),
    )

    print(
        "Acquired artifacts:",
        len(output),
    )

    print("")

    for name in (
        "FULL_20",
        "PARTIAL_14",
        "EXCLUDED_WRONG_DATASET",
    ):
        print(
            name,
            classifications[name],
        )

    accepted = sum(1 for row in output if row["validation_status"] == "ACCEPT")

    excluded = sum(1 for row in output if row["validation_status"] == "EXCLUDE")

    print("")
    print(
        "Valid PKRV observations:",
        accepted,
    )
    print(
        "Excluded acquired artifacts:",
        excluded,
    )

    print("")
    print(
        "Manifest:",
        OUTPUT_PATH.relative_to(PROJECT_ROOT),
    )

    print("")
    print("PASS: final PKRV structural acceptance contract satisfied")


if __name__ == "__main__":
    main()
