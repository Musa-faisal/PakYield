"""Build deterministic long-form normalized PKRV observations."""

from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from pakyield.config import (
    PKRV_CANONICAL_TENORS,
    PROJECT_ROOT,
    TENOR_TO_YEARS,
)
from pakyield.validation.pkrv_raw import (
    decode_text_bytes,
    extract_canonical_tenors,
)

STRUCTURAL_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_structural_acceptance.csv"

RAW_DIR = PROJECT_ROOT / "data" / "raw" / "mufap" / "pkrv"

OUTPUT_PATH = PROJECT_ROOT / "data" / "interim" / "pkrv_normalized_long.csv"

EXPECTED_DATES = 1149
EXPECTED_FULL_20 = 1119
EXPECTED_PARTIAL_14 = 30
EXPECTED_ROWS = 22800

FIRST_DATE = "2022-01-04"
LAST_DATE = "2026-09-28"

PARTIAL_MISSING = {
    "1M",
    "2M",
    "3M",
    "4M",
    "6M",
    "9M",
}

CANONICAL_SET = set(PKRV_CANONICAL_TENORS)

EXPECTED_TENOR_COUNTS = {
    tenor: (EXPECTED_FULL_20 if tenor in PARTIAL_MISSING else EXPECTED_DATES)
    for tenor in PKRV_CANONICAL_TENORS
}

OUTPUT_FIELDS = [
    "observation_date",
    "tenor",
    "tenor_years",
    "yield_pct",
    "source_organization",
    "source_file",
    "structural_class",
    "source_layout",
    "source_value_field",
]

SOURCE_ORGANIZATION = "Mutual Funds Association of Pakistan"


# Four accepted source artifacts contain exact repeated rows.
# These anomalies were exhaustively audited before normalization.
#
# Counts below represent EXTRA copies beyond the first row.
EXPECTED_EXACT_DUPLICATE_ROWS = {
    "PKRV161120222132.csv": {
        "1W": 3,
    },
    "PKRV281220222162.csv": {
        "1Y": 1,
    },
    "PKRV291220222164.csv": {
        "4M": 1,
    },
    "PKRV301220222166.csv": {
        "8Y": 1,
    },
}


def read_csv(
    path: Path,
) -> list[dict[str, str]]:
    """Read one UTF-8 CSV."""

    with path.open(
        newline="",
        encoding="utf-8",
    ) as file:
        return list(csv.DictReader(file))


def parse_decimal(
    value: object,
    *,
    context: str,
) -> Decimal:
    """Parse one source yield and reject blank/non-finite values."""

    text = str(value).strip()

    if not text:
        raise RuntimeError(f"Blank yield value: {context}")

    try:
        result = Decimal(text)

    except InvalidOperation as exc:
        raise RuntimeError(f"Invalid numeric yield {text!r}: {context}") from exc

    if not result.is_finite():
        raise RuntimeError(f"Non-finite yield {text!r}: {context}")

    if not (Decimal("0") < result < Decimal("100")):
        raise RuntimeError(f"Implausible PKRV yield {result}: {context}")

    return result


def decimal_text(
    value: Decimal,
) -> str:
    """Render Decimal without scientific notation or meaningless zeros."""

    text = format(
        value,
        "f",
    )

    if "." in text:
        text = text.rstrip("0").rstrip(".")

    return text


def exact_tenor(
    value: object,
) -> str | None:
    """Return one exact canonical tenor represented by a source label."""

    text = str(value).strip()

    if not text:
        return None

    found = extract_canonical_tenors(text)

    if len(found) == 1:
        return next(iter(found))

    return None


def locate_tenor_cell(
    values: list[str],
) -> tuple[int, str] | None:
    """Find the unique canonical-tenor cell in one CSV row."""

    matches: list[tuple[int, str]] = []

    for index, value in enumerate(values):
        tenor = exact_tenor(value)

        if tenor is not None:
            matches.append(
                (
                    index,
                    tenor,
                )
            )

    if not matches:
        return None

    if len(matches) != 1:
        raise RuntimeError(f"Ambiguous tenor row: {values!r}")

    return matches[0]


def normalized_header(
    value: object,
) -> str:
    """Normalize source header text for conservative matching."""

    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(value).strip().lower(),
    )


def find_csv_header_and_value_column(
    rows: list[list[str]],
    *,
    file_name: str,
) -> tuple[int, int, str]:
    """Find the published benchmark-value column conservatively."""

    for row_index, values in enumerate(rows):
        headers = [normalized_header(value) for value in values]

        # Direct tenor/value tables.
        if "tenor" in headers:
            if "midrate" in headers:
                return (
                    row_index,
                    headers.index("midrate"),
                    "Mid Rate",
                )

            if "avgrate" in headers:
                return (
                    row_index,
                    headers.index("avgrate"),
                    "Avg Rate",
                )

            # Some exports may split "Avg Rate"
            # across two adjacent cells.
            if "avg" in headers and "rate" in headers:
                avg_index = headers.index("avg")

                rate_index = headers.index("rate")

                if rate_index != avg_index + 1:
                    raise RuntimeError(
                        f"Ambiguous split Avg/Rate header in {file_name}: {values!r}"
                    )

                return (
                    row_index,
                    rate_index,
                    "Avg Rate",
                )

            if "rate" in headers:
                return (
                    row_index,
                    headers.index("rate"),
                    "Rate",
                )

            # Historical direct two-column format:
            #
            # Tenor,15-Aug-23
            nonblank = [value for value in values if str(value).strip()]

            if len(nonblank) == 2 and headers[0] == "tenor" and len(values) >= 2:
                return (
                    row_index,
                    1,
                    "Published tenor value",
                )

        # Dealer-table formats containing a published
        # aggregate PKRV benchmark.
        if "fmap" in headers:
            return (
                row_index,
                headers.index("fmap"),
                "FMAP",
            )

        if "avgrate" in headers:
            return (
                row_index,
                headers.index("avgrate"),
                "Avg Rate",
            )

        # Historical comma exports sometimes represent
        # the heading "Avg Rate" as:
        #
        # ...,Avg,Rate
        #
        # The aggregate benchmark value is stored in
        # the terminal Rate-position column.
        if "avg" in headers and "rate" in headers:
            avg_index = headers.index("avg")

            rate_index = headers.index("rate")

            if rate_index != avg_index + 1:
                raise RuntimeError(
                    f"Ambiguous split Avg/Rate header in {file_name}: {values!r}"
                )

            return (
                row_index,
                rate_index,
                "Avg Rate",
            )

        if "avg" in headers:
            return (
                row_index,
                headers.index("avg"),
                "Avg",
            )

    raise RuntimeError(f"No deterministic PKRV value column found in {file_name}")


def parse_csv_artifact(
    path: Path,
) -> tuple[
    dict[str, Decimal],
    str,
    str,
]:
    """Parse one comma-delimited PKRV artifact."""

    text, _ = decode_text_bytes(path.read_bytes())

    lines = [line for line in text.splitlines() if line.strip()]

    rows = list(csv.reader(lines))

    if not rows:
        raise RuntimeError(f"Empty PKRV source: {path.name}")

    # One-field rows indicate fixed-width text rather
    # than a true comma-delimited table.
    if max(len(row) for row in rows) == 1:
        return parse_fixed_width_artifact(
            text,
            file_name=path.name,
        )

    (
        header_index,
        value_column,
        value_field,
    ) = find_csv_header_and_value_column(
        rows,
        file_name=path.name,
    )

    result: dict[str, Decimal] = {}

    source_rows: dict[str, tuple[str, ...]] = {}

    duplicate_extra_rows: Counter[str] = Counter()

    for values in rows[header_index + 1 :]:
        located = locate_tenor_cell(values)

        if located is None:
            continue

        _, tenor = located

        if value_column >= len(values):
            raise RuntimeError(f"Benchmark column missing in {path.name} for {tenor}")

        parsed_yield = parse_decimal(
            values[value_column],
            context=(f"{path.name} {tenor} {value_field}"),
        )

        normalized_source_row = tuple(str(value).strip() for value in values)

        if tenor in result:
            previous_source_row = source_rows[tenor]

            previous_yield = result[tenor]

            # We collapse a repeated tenor only when
            # the entire parsed source row is identical.
            if normalized_source_row != previous_source_row:
                raise RuntimeError(
                    "Conflicting duplicate tenor row in "
                    f"{path.name}: {tenor}; "
                    f"first={previous_source_row!r}; "
                    f"duplicate={normalized_source_row!r}"
                )

            if parsed_yield != previous_yield:
                raise RuntimeError(
                    "Conflicting duplicate tenor yield in "
                    f"{path.name}: {tenor}; "
                    f"{previous_yield} != {parsed_yield}"
                )

            duplicate_extra_rows[tenor] += 1

            continue

        result[tenor] = parsed_yield

        source_rows[tenor] = normalized_source_row

    expected_duplicates = EXPECTED_EXACT_DUPLICATE_ROWS.get(
        path.name,
        {},
    )

    actual_duplicates = dict(duplicate_extra_rows)

    # Fail closed if a known anomaly changes or a new
    # duplicate anomaly appears anywhere else.
    if actual_duplicates != expected_duplicates:
        raise RuntimeError(
            "Exact-duplicate source-row contract changed "
            f"for {path.name}: "
            f"actual={actual_duplicates}; "
            f"expected={expected_duplicates}"
        )

    layout = (
        "CSV_DIRECT_RATE"
        if value_field
        in {
            "Mid Rate",
            "Rate",
            "Published tenor value",
        }
        else "CSV_PUBLISHED_BENCHMARK"
    )

    return (
        result,
        layout,
        value_field,
    )


def parse_fixed_width_artifact(
    text: str,
    *,
    file_name: str,
) -> tuple[
    dict[str, Decimal],
    str,
    str,
]:
    """Parse fixed-width historical FMA PKRV tables."""

    lines = [line.strip() for line in text.splitlines() if line.strip()]

    upper = "\n".join(lines).upper()

    fmap_with_date = "FMAP" in upper and "DATE" in upper and "TIME" in upper

    avg_rate = "AVG RATE" in upper

    if not (fmap_with_date or avg_rate):
        raise RuntimeError(f"Unsupported fixed-width PKRV layout: {file_name}")

    result: dict[str, Decimal] = {}

    for line in lines:
        pieces = line.split()

        if not pieces:
            continue

        tenor = exact_tenor(pieces[0])

        if tenor is None:
            continue

        if tenor in result:
            raise RuntimeError(f"Duplicate tenor in {file_name}: {tenor}")

        if fmap_with_date:
            # Historical FMA rows terminate with:
            # ... FMAP DATE TIME
            if len(pieces) < 4:
                raise RuntimeError(
                    f"Malformed FMAP fixed-width row: {file_name} {line!r}"
                )

            value = pieces[-3]
            value_field = "FMAP"
            layout = "FIXED_WIDTH_FMAP"

        else:
            # Older Refinitiv/FMA rows terminate
            # directly in the published Avg Rate.
            value = pieces[-1]
            value_field = "Avg Rate"
            layout = "FIXED_WIDTH_AVG_RATE"

        result[tenor] = parse_decimal(
            value,
            context=(f"{file_name} {tenor} {value_field}"),
        )

    return (
        result,
        layout,
        value_field,
    )


def parse_xlsx_artifact(
    path: Path,
) -> tuple[
    dict[str, Decimal],
    str,
    str,
]:
    """Parse one true XLSX PKRV workbook."""

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        sheets = workbook.worksheets

        if len(sheets) != 1:
            raise RuntimeError(f"Expected one PKRV worksheet in {path.name}")

        rows = [
            list(values)
            for values in sheets[0].iter_rows(values_only=True)
            if any(value is not None and str(value).strip() for value in values)
        ]

    finally:
        workbook.close()

    if not rows:
        raise RuntimeError(f"Empty PKRV XLSX: {path.name}")

    header = [normalized_header(value) for value in rows[0]]

    if "tenor" not in header or "midrate" not in header:
        raise RuntimeError(f"Unexpected PKRV XLSX header in {path.name}: {rows[0]!r}")

    tenor_column = header.index("tenor")

    value_column = header.index("midrate")

    result: dict[str, Decimal] = {}

    for values in rows[1:]:
        if tenor_column >= len(values):
            continue

        tenor = exact_tenor(values[tenor_column])

        if tenor is None:
            continue

        if tenor in result:
            raise RuntimeError(f"Duplicate tenor in {path.name}: {tenor}")

        if value_column >= len(values):
            raise RuntimeError(f"Mid Rate missing in {path.name}: {tenor}")

        result[tenor] = parse_decimal(
            values[value_column],
            context=(f"{path.name} {tenor} Mid Rate"),
        )

    return (
        result,
        "XLSX_MID_RATE",
        "Mid Rate",
    )


def parse_xml_spreadsheet_artifact(
    path: Path,
) -> tuple[
    dict[str, Decimal],
    str,
    str,
]:
    """Parse SpreadsheetML exported with an .xls extension."""

    text, detected = decode_text_bytes(path.read_bytes())

    if detected != "utf-8":
        raise RuntimeError(f"Unexpected SpreadsheetML encoding: {path.name} {detected}")

    root = ET.fromstring(text)

    namespace = {
        "ss": ("urn:schemas-microsoft-com:office:spreadsheet"),
    }

    xml_rows = root.findall(
        ".//ss:Worksheet/ss:Table/ss:Row",
        namespace,
    )

    rows: list[list[str | None]] = []

    for row in xml_rows:
        values: list[str | None] = []

        for cell in row.findall(
            "ss:Cell",
            namespace,
        ):
            data = cell.find(
                "ss:Data",
                namespace,
            )

            values.append(None if data is None else data.text)

        if any(value is not None and str(value).strip() for value in values):
            rows.append(values)

    if not rows:
        raise RuntimeError(f"Empty SpreadsheetML PKRV artifact: {path.name}")

    header = [normalized_header(value) for value in rows[0]]

    if "tenor" not in header or "midrate" not in header:
        raise RuntimeError(
            f"Unexpected SpreadsheetML PKRV header: {path.name} {rows[0]!r}"
        )

    tenor_column = header.index("tenor")

    value_column = header.index("midrate")

    result: dict[str, Decimal] = {}

    for values in rows[1:]:
        tenor = exact_tenor(values[tenor_column])

        if tenor is None:
            continue

        if tenor in result:
            raise RuntimeError(f"Duplicate tenor in {path.name}: {tenor}")

        result[tenor] = parse_decimal(
            values[value_column],
            context=(f"{path.name} {tenor} Mid Rate"),
        )

    return (
        result,
        "XML_SPREADSHEETML",
        "Mid Rate",
    )


def parse_artifact(
    path: Path,
) -> tuple[
    dict[str, Decimal],
    str,
    str,
]:
    """Dispatch one accepted raw artifact to its deterministic parser."""

    suffix = path.suffix.lower()

    if suffix == ".xlsx":
        return parse_xlsx_artifact(path)

    if suffix == ".xls":
        return parse_xml_spreadsheet_artifact(path)

    if suffix == ".csv":
        return parse_csv_artifact(path)

    raise RuntimeError(f"Unsupported accepted PKRV extension: {path.name}")


def expected_tenors_for_manifest_row(
    row: dict[str, str],
) -> set[str]:
    """Return the exact tenor universe permitted for one accepted date."""

    structural_class = row["structural_class"]

    if structural_class == "FULL_20":
        return set(CANONICAL_SET)

    if structural_class == "PARTIAL_14":
        return set(CANONICAL_SET) - PARTIAL_MISSING

    raise RuntimeError(f"Unexpected accepted structural class: {structural_class}")


def main() -> None:
    """Build and validate the full long-form PKRV normalization artifact."""

    manifest = read_csv(STRUCTURAL_PATH)

    accepted = [row for row in manifest if row["validation_status"] == "ACCEPT"]

    if len(accepted) != EXPECTED_DATES:
        raise RuntimeError(f"Expected 1,149 accepted PKRV dates; found {len(accepted)}")

    accepted.sort(key=lambda row: row["observation_date"])

    if accepted[0]["observation_date"] != FIRST_DATE:
        raise RuntimeError("Unexpected first PKRV date")

    if accepted[-1]["observation_date"] != LAST_DATE:
        raise RuntimeError("Unexpected final PKRV date")

    class_counts = Counter(row["structural_class"] for row in accepted)

    if dict(class_counts) != {
        "FULL_20": EXPECTED_FULL_20,
        "PARTIAL_14": EXPECTED_PARTIAL_14,
    }:
        raise RuntimeError(
            f"Accepted structural-class counts changed: {dict(class_counts)}"
        )

    output: list[dict[str, str]] = []

    layout_counts = Counter()

    date_row_counts = Counter()

    for row in accepted:
        observation_date = row["observation_date"]

        file_name = row["file_name"]

        path = RAW_DIR / file_name

        if not path.is_file():
            raise RuntimeError(f"Accepted PKRV artifact missing: {file_name}")

        (
            values,
            layout,
            value_field,
        ) = parse_artifact(path)

        expected_tenors = expected_tenors_for_manifest_row(row)

        actual_tenors = set(values)

        if actual_tenors != expected_tenors:
            raise RuntimeError(
                "Normalized tenor set disagrees with "
                "structural contract: "
                f"{observation_date} {file_name} "
                f"expected={sorted(expected_tenors)} "
                f"actual={sorted(actual_tenors)}"
            )

        expected_count = int(row["tenor_count"])

        if len(values) != expected_count:
            raise RuntimeError(
                "Normalized tenor count disagrees with "
                "structural manifest: "
                f"{file_name} "
                f"{len(values)} != {expected_count}"
            )

        layout_counts[layout] += 1

        for tenor in PKRV_CANONICAL_TENORS:
            if tenor not in values:
                continue

            output.append(
                {
                    "observation_date": (observation_date),
                    "tenor": tenor,
                    "tenor_years": format(
                        TENOR_TO_YEARS[tenor],
                        ".12g",
                    ),
                    "yield_pct": decimal_text(values[tenor]),
                    "source_organization": (SOURCE_ORGANIZATION),
                    "source_file": file_name,
                    "structural_class": row["structural_class"],
                    "source_layout": (layout),
                    "source_value_field": (value_field),
                }
            )

            date_row_counts[observation_date] += 1

    if len(output) != EXPECTED_ROWS:
        raise RuntimeError(
            f"Unexpected normalized PKRV row count: {len(output)} != {EXPECTED_ROWS}"
        )

    keys = {
        (
            row["observation_date"],
            row["tenor"],
        )
        for row in output
    }

    if len(keys) != EXPECTED_ROWS:
        raise RuntimeError("Duplicate normalized PKRV date-tenor keys detected")

    tenor_counts = Counter(row["tenor"] for row in output)

    if dict(tenor_counts) != (EXPECTED_TENOR_COUNTS):
        raise RuntimeError(
            f"Normalized PKRV tenor counts changed: {dict(tenor_counts)}"
        )

    observed_date_counts = Counter(date_row_counts.values())

    if dict(observed_date_counts) != {
        20: EXPECTED_FULL_20,
        14: EXPECTED_PARTIAL_14,
    }:
        raise RuntimeError(
            f"Per-date PKRV row counts changed: {dict(observed_date_counts)}"
        )

    by_key = {
        (
            row["observation_date"],
            row["tenor"],
        ): row["yield_pct"]
        for row in output
    }

    expected_sentinels = {
        (
            "2022-01-04",
            "1W",
        ): "9.99",
        (
            "2022-01-04",
            "20Y",
        ): "12.38",
        (
            "2022-04-28",
            "1W",
        ): "12.58",
        (
            "2022-04-28",
            "1M",
        ): "13.05",
        (
            "2022-04-28",
            "20Y",
        ): "12.97",
        (
            "2023-06-16",
            "1W",
        ): "21.53",
        (
            "2023-06-16",
            "2W",
        ): "21.43",
        (
            "2023-06-16",
            "1Y",
        ): "21.98",
        (
            "2025-08-20",
            "1W",
        ): "11",
        (
            "2025-08-20",
            "20Y",
        ): "12.44",
        (
            "2026-03-18",
            "1W",
        ): "10.5",
        (
            "2026-03-18",
            "20Y",
        ): "12.71",
        (
            "2026-09-28",
            "1W",
        ): "11.54",
        (
            "2026-09-28",
            "20Y",
        ): "12.9",
    }

    for key, expected in expected_sentinels.items():
        actual = by_key.get(key)

        if actual != expected:
            raise RuntimeError(
                f"PKRV sentinel mismatch {key}: {actual!r} != {expected!r}"
            )

    # The six monthly tenors must remain absent
    # on all 30 PARTIAL_14 dates.
    partial_dates = {
        row["observation_date"]
        for row in accepted
        if row["structural_class"] == "PARTIAL_14"
    }

    for observation_date in partial_dates:
        for tenor in PARTIAL_MISSING:
            if (
                observation_date,
                tenor,
            ) in keys:
                raise RuntimeError(
                    f"Forbidden manufactured tenor detected: {observation_date} {tenor}"
                )

    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary = OUTPUT_PATH.with_suffix(OUTPUT_PATH.suffix + ".part")

    with temporary.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=OUTPUT_FIELDS,
        )

        writer.writeheader()

        writer.writerows(output)

    temporary.replace(OUTPUT_PATH)

    year_counts = Counter(row["observation_date"][:4] for row in output)

    print("===== PKRV NORMALIZED LONG FORM =====")

    print(
        "Accepted dates:",
        len(accepted),
    )

    print(
        "FULL_20 dates:",
        class_counts["FULL_20"],
    )

    print(
        "PARTIAL_14 dates:",
        class_counts["PARTIAL_14"],
    )

    print(
        "Normalized rows:",
        len(output),
    )

    print(
        "Unique date-tenor keys:",
        len(keys),
    )

    print(
        "Date row-count distribution:",
        dict(sorted(observed_date_counts.items())),
    )

    print(
        "Rows by year:",
        dict(sorted(year_counts.items())),
    )

    print("")
    print("Source layouts:")

    for layout, count in sorted(layout_counts.items()):
        print(
            " ",
            layout,
            count,
        )

    print("")
    print("Tenor counts:")

    for tenor in PKRV_CANONICAL_TENORS:
        print(
            " ",
            tenor,
            tenor_counts[tenor],
        )

    print("")
    print(
        "Output:",
        OUTPUT_PATH.relative_to(PROJECT_ROOT),
    )

    print("")
    print("PASS: deterministic PKRV long-form normalization contract satisfied")


if __name__ == "__main__":
    main()
