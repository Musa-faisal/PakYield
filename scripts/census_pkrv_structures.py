"""Census structural tenor coverage across acquired PKRV source artifacts."""

from __future__ import annotations

import csv
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from pakyield.config import PKRV_CANONICAL_TENORS

PROJECT_ROOT = Path(__file__).resolve().parents[1]

RAW_DIR = PROJECT_ROOT / "data" / "raw" / "mufap" / "pkrv"

OUTPUT_PATH = PROJECT_ROOT / "data" / "manifests" / "pkrv_structural_census.csv"

FIELDS = [
    "file_name",
    "file_extension",
    "file_size_bytes",
    "structure_family",
    "canonical_tenor_count",
    "full_20_tenors",
    "missing_tenors",
    "dataset_hints",
    "first_nonempty_row",
]

CANONICAL = tuple(tenor.upper() for tenor in PKRV_CANONICAL_TENORS)

CANONICAL_SET = set(CANONICAL)


def normalize(value: object) -> str:
    """Normalize a source cell for structural matching."""

    if value is None:
        return ""

    return str(value).strip()


def tenor_coverage(
    rows: list[list[str]],
) -> tuple[set[str], list[str]]:
    """Return discovered canonical tenors and flattened text."""

    found: set[str] = set()
    flattened: list[str] = []

    for row in rows:
        for cell in row:
            value = normalize(cell)
            upper = value.upper()

            if upper:
                flattened.append(upper)

            if upper in CANONICAL_SET:
                found.add(upper)

    return found, flattened


def classify_hints(
    file_name: str,
    flattened: list[str],
) -> list[str]:
    """Return conservative dataset/content hints."""

    joined = " | ".join(flattened)

    hints: list[str] = []

    upper_name = file_name.upper()

    if upper_name.startswith("PKFRV"):
        hints.append("PKFRV_FILENAME")

    if upper_name.startswith("PKISRV"):
        hints.append("PKISRV_FILENAME")

    if "KIBOR FIXING" in joined:
        hints.append("KIBOR_CONTENT")

    if "PIB-FRB-" in joined:
        hints.append("FLOATING_RATE_PIB_CONTENT")

    if "PKISRV" in joined or "IJARAH" in joined:
        hints.append("ISLAMIC_CONTENT")

    if "PKRV" in joined:
        hints.append("PKRV_CONTENT")

    if (
        "REVALUATION RATES" in joined
        or "REVALUATION | RATES" in joined
        or "T-BILL & PIB" in joined
        or "T-BILLS/PIBS" in joined
    ):
        hints.append("SOVEREIGN_REVALUATION_CONTENT")

    return hints


def csv_rows(
    path: Path,
) -> list[list[str]]:
    """Read CSV source rows without altering source bytes."""

    raw = path.read_bytes()

    text = raw.decode(
        "utf-8-sig",
        errors="replace",
    )

    return [[normalize(cell) for cell in row] for row in csv.reader(text.splitlines())]


def xlsx_rows(
    path: Path,
) -> list[list[str]]:
    """Read XLSX workbook rows."""

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook[workbook.sheetnames[0]]

    rows = [
        [normalize(cell) for cell in row] for row in sheet.iter_rows(values_only=True)
    ]

    workbook.close()

    return rows


def xml_xls_rows(
    path: Path,
) -> list[list[str]]:
    """Read Spreadsheet XML stored with an .xls extension."""

    root = ET.fromstring(path.read_bytes())

    rows: list[list[str]] = []

    for element in root.iter():
        if element.tag.split("}")[-1] != "Row":
            continue

        values: list[str] = []

        for cell in element:
            if cell.tag.split("}")[-1] != "Cell":
                continue

            value = ""

            for child in cell:
                if child.tag.split("}")[-1] == "Data":
                    value = normalize(child.text)
                    break

            values.append(value)

        rows.append(values)

    return rows


def nonempty_rows(
    rows: list[list[str]],
) -> list[list[str]]:
    """Remove fully blank rows for structural reporting only."""

    return [row for row in rows if any(normalize(cell) for cell in row)]


def structure_family(
    rows: list[list[str]],
    tenor_count: int,
) -> str:
    """Assign a descriptive source structure family."""

    clean = nonempty_rows(rows)

    if not clean:
        return "EMPTY"

    first = tuple(normalize(cell) for cell in clean[0])

    first_three = first[:3]

    if first_three == (
        "Tenor",
        "Mid Rate",
        "Change",
    ):
        return "CANONICAL_TENOR_MID_CHANGE"

    if tenor_count == 20:
        return "LEGACY_FULL_20_TENOR"

    if tenor_count > 0:
        return "PARTIAL_CANONICAL_TENORS"

    return "NO_CANONICAL_TENORS"


def main() -> None:
    """Run structural census across all acquired artifacts."""

    files = sorted(
        path
        for path in RAW_DIR.iterdir()
        if path.is_file()
        and path.name != ".gitkeep"
        and not path.name.endswith(".part")
    )

    if len(files) != 1157:
        raise RuntimeError(f"Expected 1,157 acquired artifacts; found {len(files)}")

    output: list[dict[str, str]] = []

    for path in files:
        suffix = path.suffix.lower()

        if suffix == ".csv":
            rows = csv_rows(path)
        elif suffix == ".xlsx":
            rows = xlsx_rows(path)
        elif suffix == ".xls":
            rows = xml_xls_rows(path)
        else:
            raise RuntimeError(f"Unsupported source format: {path}")

        clean = nonempty_rows(rows)

        found, flattened = tenor_coverage(rows)

        missing = [tenor for tenor in CANONICAL if tenor not in found]

        hints = classify_hints(
            path.name,
            flattened,
        )

        family = structure_family(
            rows,
            len(found),
        )

        first_row = " | ".join(clean[0][:12]) if clean else ""

        output.append(
            {
                "file_name": path.name,
                "file_extension": suffix,
                "file_size_bytes": str(path.stat().st_size),
                "structure_family": family,
                "canonical_tenor_count": str(len(found)),
                "full_20_tenors": ("TRUE" if len(found) == 20 else "FALSE"),
                "missing_tenors": "|".join(missing),
                "dataset_hints": "|".join(hints),
                "first_nonempty_row": (first_row),
            }
        )

    temporary = OUTPUT_PATH.with_suffix(OUTPUT_PATH.suffix + ".part")

    with temporary.open(
        "w",
        newline="",
        encoding="utf-8",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=FIELDS,
        )
        writer.writeheader()
        writer.writerows(output)

    temporary.replace(OUTPUT_PATH)

    families = Counter(row["structure_family"] for row in output)

    coverage = Counter(int(row["canonical_tenor_count"]) for row in output)

    print("===== PKRV STRUCTURAL CENSUS =====")
    print("Artifacts:", len(output))

    print("")
    print("===== STRUCTURE FAMILIES =====")

    for family, count in sorted(families.items()):
        print(
            family,
            count,
        )

    print("")
    print("===== TENOR COVERAGE =====")

    for count in sorted(
        coverage,
        reverse=True,
    ):
        print(
            f"{count:2d} tenors:",
            coverage[count],
        )

    suspicious = [
        row
        for row in output
        if (
            row["full_20_tenors"] != "TRUE"
            or row["dataset_hints"].find("PKFRV") >= 0
            or row["dataset_hints"].find("KIBOR_CONTENT") >= 0
            or row["dataset_hints"].find("ISLAMIC_CONTENT") >= 0
        )
    ]

    print("")
    print(
        "Files requiring review:",
        len(suspicious),
    )

    print("")
    print("===== REVIEW FILES =====")

    for row in suspicious:
        print(
            row["file_name"],
            "|",
            row["canonical_tenor_count"],
            "tenors",
            "|",
            row["structure_family"],
            "|",
            row["dataset_hints"],
        )

    print("")
    print(
        "Census:",
        OUTPUT_PATH.relative_to(PROJECT_ROOT),
    )


if __name__ == "__main__":
    main()
